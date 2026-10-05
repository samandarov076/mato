"""08-hisobotlar: savdo, tushum, foyda, top tovarlar, qoldiq qiymati, nasiya, narx og'ishi, Excel."""
import uuid
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, update

from app.core.db import system_session
from app.modules.credit.models import Debt
from app.modules.sales.models import Sale
from tests.shop import TODAY, TODAY_DATE

YESTERDAY = (TODAY_DATE - timedelta(days=1)).isoformat()
ZEROS = {"sales": 0, "returns": 0, "profit": 0, "cash": 0, "card": 0, "count": 0, "meters": "0"}


async def _price(s, batch: int, price: int) -> None:
    card = (await s.client.get(f"/api/batches/{batch}", headers=s.headers)).json()
    await s.client.patch(
        f"/api/catalog/skus/{card['sku']['id']}", json={"price": price, "price_unit": "m"}, headers=s.headers
    )


@pytest.fixture
async def trade(shop, seller):
    """Bir kunlik savdo. Tannarx: A — 30 000, B — 20 000 so'm/m; tavsiya narx: A — 45 000, B — 30 000.

    Kecha:  A 2 m × 45 000 = 90 000 (naqd).
    Bugun:  1) A 10 m × 44 000 = 440 000 naqd; shundan 2 m qaytdi (88 000, naqd)
            2) B 5 m × 30 000 = 150 000 karta (sotuvchi)
            3) A 4 m × 50 000 = 200 000 nasiya, boshlang'ich to'lov 100 000 karta → qarz 100 000
            4) A 1 m — bekor qilingan (hisobga kirmaydi)
            qarz to'lovi 40 000 naqd, hisobdan chiqarildi 10 000 → qarz 50 000
    """
    s = await shop()
    a = await s.in_store(0, "200", cost_per_m=30000)
    b = await s.in_store(1, "100", 2, cost_per_m=20000)
    await _price(s, a, 45000)
    await _price(s, b, 30000)
    nodira = await seller(s)
    customer = await s.customer()

    old = (await s.sell_to((a, "2", 45000))).json()["id"]
    async with system_session() as db:
        await db.execute(
            update(Sale).where(Sale.id == old).values(created_at=datetime.now(UTC) - timedelta(days=1))
        )
    first = (await s.sell_to((a, "10", 44000))).json()["id"]
    await s.sell_to((b, "5", 30000), headers=nodira, payment_type="card")
    credit = await s.sell_to(
        (a, "4", 50000), payment_type="credit", customer_id=customer,
        initial_payment=100_000, initial_payment_method="card",
    )  # fmt: skip
    cancelled = (await s.sell_to((a, "1", 45000))).json()["id"]
    await s.client.post(f"/api/sales/{cancelled}/cancel", json={"reason": "Xato"}, headers=s.headers)
    line = (await s.sale(first))["lines"][0]["id"]
    await s.client.post(
        f"/api/sales/{first}/returns", json={"lines": [{"sale_line_id": line, "meters": "2"}]}, headers=s.headers
    )
    await s.client.post(
        f"/api/customers/{customer}/payments", json={"client_uuid": str(uuid.uuid4()), "amount": 40_000, "method": "cash"}, headers=s.headers
    )
    async with system_session() as db:
        debt = await db.scalar(select(Debt.id).where(Debt.sale_id == credit.json()["id"]))
    await s.client.post(
        f"/api/debts/{debt}/writeoff", json={"amount": 10_000, "reason": "Chegirma"}, headers=s.headers
    )
    s.batches, s.seller_headers = (a, b), nodira
    return s


async def _get(s, path: str, headers=None, **params):
    return await s.client.get(f"/api/reports/{path}", params=params, headers=headers or s.headers)


def _today(**params) -> dict:
    return {"date_from": TODAY, "date_to": TODAY, **params}


# ====================== Savdo, tushum, foyda ======================

TODAY_MONEY = {
    "sales": 702_000,  # 440 000 + 150 000 + 200 000 − 88 000 qaytarilgan
    "returns": 88_000,
    "profit": 242_000,  # (140 000 + 50 000 + 80 000) − qaytarilgan 2 m foydasi 28 000
    "cash": 392_000,  # 440 000 + qarz to'lovi 40 000 − qaytarilgan 88 000
    "card": 250_000,  # 150 000 + boshlang'ich to'lov 100 000
    "count": 3,
    "meters": "17.00",
}
YESTERDAY_MONEY = {
    "sales": 90_000, "returns": 0, "profit": 30_000, "cash": 90_000, "card": 0, "count": 1, "meters": "2.00",
}  # fmt: skip


async def test_sales_report_by_day(trade):
    report = (await _get(trade, "sales", date_from=YESTERDAY, date_to=TODAY)).json()

    assert report["group"] == "day"
    assert report["rows"] == [{"period": YESTERDAY, **YESTERDAY_MONEY}, {"period": TODAY, **TODAY_MONEY}]
    assert report["total"] == {
        "sales": 792_000, "returns": 88_000, "profit": 272_000, "cash": 482_000, "card": 250_000,
        "count": 4, "meters": "19.00",
    }  # fmt: skip
    assert report["previous"] == ZEROS  # undan oldingi ikki kunda savdo bo'lmagan

    # Bugunni kecha bilan taqqoslash
    today = (await _get(trade, "sales", **_today())).json()
    assert (today["total"], today["previous"]) == (TODAY_MONEY, YESTERDAY_MONEY)

    months = (await _get(trade, "sales", date_from=YESTERDAY, date_to=TODAY, group="month")).json()
    assert months["rows"][-1]["period"] == TODAY[:7]
    assert sum(row["sales"] for row in months["rows"]) == 792_000


async def test_sales_report_filters(trade, new_location):
    s = trade

    async def total(**params) -> dict:
        return (await _get(s, "sales", **_today(**params))).json()["total"]

    assert await total(store_id=s.store) == TODAY_MONEY
    other = await new_location(s.business_id, "store", name="Boshqa do'kon")
    assert await total(store_id=other) == ZEROS

    # Sotuvchi faqat B ni kartaga sotgan
    sellers = (await s.client.get("/api/users", headers=s.headers)).json()
    nodira = next(user["id"] for user in sellers if user["role"] == "seller")
    assert await total(seller_id=nodira) == {
        "sales": 150_000, "returns": 0, "profit": 50_000, "cash": 0, "card": 150_000, "count": 1, "meters": "5.00",
    }  # fmt: skip

    # Mato bo'yicha: savdo va foyda bor, tushum hisoblanmaydi (pul matoga bog'lanmaydi)
    by_fabric = await total(fabric_id=s.fabric)
    assert (by_fabric["sales"], by_fabric["profit"], by_fabric["cash"], by_fabric["card"]) == (
        702_000, 242_000, None, None,
    )  # fmt: skip


async def test_report_period_is_validated(trade):
    backwards = await _get(trade, "sales", date_from=TODAY, date_to=YESTERDAY)
    assert backwards.json()["fields"] == {"date_to": "invalid_range"}
    too_long = await _get(trade, "sales", date_from="2024-01-01", date_to=TODAY)
    assert too_long.status_code == 422
    assert (await _get(trade, "sales")).status_code == 422  # davr majburiy


# ====================== Bosh sahifa ======================


async def test_dashboard(trade, new_location):
    s = trade
    card = (await s.client.get(f"/api/batches/{s.batches[0]}", headers=s.headers)).json()
    await s.client.patch(
        f"/api/catalog/skus/{card['sku']['id']}", json={"low_stock_threshold": "250"}, headers=s.headers
    )  # do'konda A dan 186 m qolgan — chegaradan past

    board = (await _get(s, "dashboard")).json()

    assert (board["date_from"], board["date_to"]) == (TODAY, TODAY)
    assert (board["current"], board["previous"]) == (TODAY_MONEY, YESTERDAY_MONEY)
    assert (board["debt_total"], board["debt_overdue"], board["low_stock"]) == (50_000, 0, 1)
    assert (board["diff_amount"], board["diff_m"]) == (0, "0")
    assert len(board["chart"]) == 30
    assert board["chart"][-2:] == [{"day": YESTERDAY, "sales": 90_000}, {"day": TODAY, "sales": 702_000}]

    week = (await _get(s, "dashboard", period="week")).json()
    assert week["date_from"] == (TODAY_DATE - timedelta(days=TODAY_DATE.weekday())).isoformat()
    month = (await _get(s, "dashboard", period="month")).json()
    assert (month["date_from"], month["date_to"]) == (TODAY_DATE.replace(day=1).isoformat(), TODAY)
    assert month["current"]["sales"] >= 702_000

    other = await new_location(s.business_id, "store", name="Boshqa do'kon")
    empty = (await _get(s, "dashboard", store_id=other)).json()
    assert (empty["current"], empty["debt_total"], empty["low_stock"]) == (ZEROS, 0, 0)
    assert {day["sales"] for day in empty["chart"]} == {0}


async def test_dashboard_counts_overdue_debt_and_meter_difference(trade):
    s = trade
    async with system_session() as db:
        await db.execute(
            update(Debt).where(Debt.business_id == s.business_id).values(due_date=TODAY_DATE - timedelta(days=1))
        )
    # B partiyasi do'konda yopildi: tizimda 95 m qolgan edi — kam chiqdi (95 × 20 000 = 1 900 000)
    await s.client.post(
        f"/api/batches/{s.batches[1]}/close", json={"location_id": s.store, "note": "Tugadi"}, headers=s.headers
    )

    board = (await _get(s, "dashboard")).json()

    assert (board["debt_total"], board["debt_overdue"]) == (50_000, 50_000)
    assert (board["diff_amount"], board["diff_m"]) == (1_900_000, "95.00")


# ====================== Soatlar ======================


async def test_hours_report(trade):
    report = (await _get(trade, "hours", **_today())).json()

    hour = datetime.now(ZoneInfo("Asia/Tashkent")).hour
    assert sum(row["count"] for row in report["hours"]) == 3
    assert sum(row["amount"] for row in report["hours"]) == 790_000  # qaytarish ayirilmaydi — sotuv payti
    assert {row["hour"] for row in report["hours"]} <= {hour, (hour - 1) % 24}  # soat almashgan bo'lishi mumkin
    weekdays = {row["weekday"]: row for row in report["weekdays"]}
    assert len(weekdays) == 7
    assert (weekdays[TODAY_DATE.isoweekday()]["count"], weekdays[TODAY_DATE.isoweekday()]["amount"]) == (3, 790_000)


# ====================== Top tovarlar ======================


async def test_top_items(trade):
    s = trade

    async def top(**params) -> list[tuple]:
        rows = (await _get(s, "top", **_today(**params))).json()
        return [(row["label"], row["meters"], row["amount"], row["profit"]) for row in rows]

    # Sof ko'rsatkichlar: A dan 14 m sotilib, 2 m qaytgan
    red = ("M101-01 · Atlas · Rang 01", "12.00", 552_000, 192_000)
    blue = ("M101-02 · Atlas · Rang 02", "5.00", 150_000, 50_000)
    assert await top() == [red, blue]
    assert await top(metric="meters", limit=1) == [red]
    assert await top(by="color") == [("01 · Rang 01", *red[1:]), ("02 · Rang 02", *blue[1:])]
    assert await top(by="fabric") == [("M101 · Atlas", "17.00", 702_000, 242_000)]
    assert [row[1:] for row in await top(by="pattern")] == [("17.00", 702_000, 242_000)]
    assert await top(date_from=YESTERDAY, date_to=YESTERDAY) == [
        ("M101-01 · Atlas · Rang 01", "2.00", 90_000, 30_000)
    ]


# ====================== Qoldiq ======================


async def test_stock_value_by_location(trade):
    s = trade
    card = (await s.client.get(f"/api/batches/{s.batches[0]}", headers=s.headers)).json()
    await s.client.patch(
        f"/api/catalog/skus/{card['sku']['id']}", json={"low_stock_threshold": "250"}, headers=s.headers
    )

    report = (await _get(s, "stock")).json()

    places = [
        (p["name"], p["type"], p["meters"], p["rolls"], p["value"], p["skus"], p["low"])
        for p in report["places"]
    ]
    assert places == [
        # Sklad: A 300 m × 30 000 + B 400 m × 20 000
        ("Asosiy sklad", "warehouse", "700.00", 14, 17_000_000, 2, 0),
        # Do'kon: A 186 m (200 − 2 − 10 − 4 + 2 qaytgan), B 95 m
        ("Chorsu do'koni", "store", "281.00", 6, 7_480_000, 2, 1),
    ]
    assert (report["meters"], report["value"]) == ("981.00", 24_480_000)


# ====================== Nasiya ======================


@contextmanager
def _selects():
    """Biznes ulanishida bajarilgan SELECT so'rovlarini sanaydi (sessiyani sozlash so'rovlarisiz)."""
    from app.core.db import engine

    seen: list[str] = []

    def record(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith("SELECT") and "set_config" not in statement:
            seen.append(statement)

    event.listen(engine.sync_engine, "before_cursor_execute", record)
    try:
        yield seen
    finally:
        event.remove(engine.sync_engine, "before_cursor_execute", record)


async def test_dashboard_and_stock_report_use_few_queries(trade, new_location):
    """Tezlik: bosh sahifa ~25 so'rov qilardi; qoldiq hisoboti har joy uchun alohida so'rov (N+1)."""
    s = trade
    with _selects() as board:
        assert (await _get(s, "dashboard", period="month")).status_code == 200
    assert len(board) <= 10, len(board)

    with _selects() as few:
        before = (await _get(s, "stock")).json()
    for index in range(5):
        await new_location(s.business_id, "store", name=f"Yangi do'kon {index}")
    with _selects() as many:
        after = (await _get(s, "stock")).json()
    assert len(after["places"]) == len(before["places"]) + 5
    assert len(many) == len(few), (len(few), len(many))  # joylar soniga bog'liq emas
    assert len(many) <= 6


async def test_credit_report(trade, new_location):
    s = trade
    report = (await _get(s, "credit", **_today())).json()
    assert report == {
        "debt_total": 50_000, "debt_overdue": 0, "debtors": 1,
        "aging": {"total": 50_000, "not_due": 50_000, "d1_30": 0, "d31_60": 0, "d60_plus": 0},
        "given": 100_000, "paid": 40_000, "returned": 0, "written_off": 10_000,
    }  # fmt: skip

    earlier = (await _get(s, "credit", date_from=YESTERDAY, date_to=YESTERDAY)).json()
    assert (earlier["debt_total"], earlier["given"], earlier["paid"]) == (50_000, 0, 0)  # qarz — hozirgi holat
    other = await new_location(s.business_id, "store", name="Boshqa do'kon")
    assert (await _get(s, "credit", **_today(store_id=other))).json()["debt_total"] == 0


# ====================== Narx og'ishi ======================


async def test_price_deviation(trade):
    s = trade
    report = (await _get(s, "price-deviation", **_today())).json()

    # A: 10 m 44 000 dan (−10 000) va 4 m 50 000 dan (+20 000), tavsiya 45 000; B — tavsiya narxda
    assert [(r["label"], r["meters"], r["expected"], r["actual"], r["diff"], r["percent"]) for r in report["rows"]] == [
        ("M101-02 · Atlas · Rang 02", "5.00", 150_000, 150_000, 0, "0.0"),
        ("M101-01 · Atlas · Rang 01", "14.00", 630_000, 640_000, 10_000, "1.6"),
    ]
    assert (report["expected"], report["actual"], report["diff"], report["percent"]) == (
        780_000, 790_000, 10_000, "1.3",
    )  # fmt: skip

    sellers = (await _get(s, "price-deviation", **_today(by="seller"))).json()
    assert [(r["label"], r["diff"]) for r in sellers["rows"]] == [("Nodira Sotuvchi", 0), ("Akmal Ega", 10_000)]


# ====================== Excel ======================


async def test_excel_export(trade):
    s = trade

    async def sheet(path: str, **params):
        response = await _get(s, path, **params)
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert ".xlsx" in response.headers["content-disposition"]
        return [list(row) for row in load_workbook(BytesIO(response.content)).worksheets[0].values]

    sales = await sheet("sales/export", date_from=YESTERDAY, date_to=TODAY)
    assert sales[0][:3] == ["Davr", "Savdo, so'm", "Qaytarish, so'm"]
    assert sales[1][:2] == [YESTERDAY, 90_000]
    assert sales[-1] == ["Jami", 792_000, 88_000, 482_000, 250_000, 272_000, 4, 19]

    top = await sheet("top/export", **_today())
    assert top[1] == ["M101-01 · Atlas · Rang 01", 12, 552_000, 192_000]
    stock = await sheet("stock/export")
    assert stock[1][:3] == ["Asosiy sklad", "Sklad", 700]
    assert stock[-1][0] == "Jami"
    credit = await sheet("credit/export")
    assert credit[1][:4] == ["Sardor Usta", "+998901234567", 50_000, 0]
    hours = await sheet("hours/export", **_today())
    assert hours[0] == ["Soat", "Sotuvlar soni", "Summa, so'm"]
    deviation = await sheet("price-deviation/export", **_today(by="seller"))
    assert deviation[0][0] == "Sotuvchi"
    assert deviation[-1][:1] + deviation[-1][2:5] == ["Jami", 780_000, 790_000, 10_000]


async def test_list_exports(trade, keeper):
    """Ro'yxatlar Excel'ga: sotuvlar, mijozlar, qoldiq (matritsa), metr farqi, jurnal — joriy filtr bilan."""
    s = trade
    dilshod = await keeper(s)

    async def book(path: str, headers=None, **params):
        response = await s.client.get(f"/api{path}", params=params, headers=headers or s.headers)
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert ".xlsx" in response.headers["content-disposition"]
        return [[list(row) for row in sheet.values] for sheet in load_workbook(BytesIO(response.content)).worksheets]

    async def sheet(path: str, headers=None, **params):
        return (await book(path, headers, **params))[0]

    # --- Sotuvlar ---
    sales = await sheet("/sales/export")
    assert sales[0] == [
        "Raqam", "Vaqt", "Do'kon", "Sotuvchi", "Mijoz", "To'lov turi", "Summa, so'm", "Qaytarilgan, so'm", "Holat",
    ]  # fmt: skip
    assert len(sales) == 6 and all(row[0].startswith("S-") for row in sales[1:])
    returned = next(row for row in sales[1:] if row[7])
    assert (returned[2], returned[5], returned[6], returned[7], returned[8]) == (
        "Chorsu do'koni", "Naqd", 440_000, 88_000, "Qisman qaytarilgan",
    )  # fmt: skip
    credit = await sheet("/sales/export", payment_type="credit")
    assert [(row[4], row[5], row[6]) for row in credit[1:]] == [("Sardor Usta", "Nasiya", 200_000)]
    assert len(await sheet("/sales/export", date_from=YESTERDAY, date_to=YESTERDAY)) == 2

    # --- Mijozlar (oxirgi to'lov summasi bilan) ---
    customers = await sheet("/customers/export")
    assert customers[0] == [
        "Mijoz", "Telefon", "Qarz, so'm", "Muddati o'tgan, so'm", "Eng eski qarz, kun",
        "Oxirgi to'lov", "Oxirgi to'lov, so'm", "Izoh",
    ]  # fmt: skip
    assert customers[1][:5] == ["Sardor Usta", "+998901234567", 50_000, 0, 0]
    assert customers[1][6] == 40_000
    listed = (await s.client.get("/api/customers", headers=s.headers)).json()["items"][0]
    assert listed["last_payment_amount"] == 40_000

    # --- Qoldiq: matritsa (tovar × joy) ---
    stock = await sheet("/stock/export")
    assert stock[0] == ["Kod", "Tovar", "Asosiy sklad", "Chorsu do'koni", "Jami, m"]
    assert stock[1] == ["M101-01", "Atlas · Rang 01", 300, 186, 486]
    assert stock[2] == ["M101-02", "Atlas · Rang 02", 400, 95, 495]
    only_store = await sheet("/stock/export", location_id=s.store)
    assert only_store[0] == ["Kod", "Tovar", "Chorsu do'koni", "Jami, m"]

    # --- Tannarx, foyda, qiymat sotuvchi va omborchi fayllarida yo'q ---
    secret = ("tannarx", "foyda", "qiymat")
    costs = {30_000, 20_000}  # partiyalar tannarxi — hech bir katakda turmasligi kerak
    for path, who in (
        ("/sales/export", s.seller_headers), ("/customers/export", s.seller_headers),
        ("/stock/export", s.seller_headers), ("/stock/export", dilshod),
    ):  # fmt: skip
        cells = [cell for rows in await book(path, who) for row in rows for cell in row]
        text = " ".join(str(cell) for cell in cells).lower()
        assert not any(word in text for word in secret), (path, text)
        assert not costs & {cell for cell in cells if isinstance(cell, int | float)}, path

    # --- Metr farqi va jurnal — faqat ega ---
    diff = await book("/stock/differences/export")
    assert diff[0][0][:4] == ["Partiya", "Tovar", "Joy", "Yozilgan metr"]
    assert diff[1][0][:3] == ["Ko'chirish", "Partiya", "Tovar"]
    journal = await sheet("/audit-logs/export", exclude_category="login")
    assert journal[0] == ["Vaqt", "Xodim", "Joy", "Kategoriya", "Hodisa", "Tafsilot"]
    assert {row[3] for row in journal[1:]} >= {"Sotuv", "Nasiya"} and "Kirish" not in {row[3] for row in journal[1:]}
    assert any(row[4] == "sale.created" and "S-0000" in row[5] for row in journal[1:])
    for path in ("/stock/differences/export", "/audit-logs/export"):
        for who in (s.seller_headers, dilshod):
            assert (await s.client.get(f"/api{path}", headers=who)).status_code == 403
    assert (await s.client.get("/api/sales/export", headers=dilshod)).status_code == 403


# ====================== Ruxsatlar ======================


async def test_reports_are_for_owners_only(trade, keeper, shop, client):
    s = trade
    dilshod = await keeper(s)
    for headers in (dilshod, s.seller_headers):
        assert (await _get(s, "dashboard", headers=headers)).status_code == 403
        assert (await _get(s, "sales", headers=headers, **_today())).status_code == 403
        assert (await _get(s, "sales/export", headers=headers, **_today())).status_code == 403
    assert (await client.get("/api/reports/dashboard")).status_code == 401

    # Boshqa biznes faqat o'z ma'lumotini ko'radi
    stranger = await shop()
    assert (await _get(stranger, "sales", **_today())).json()["total"] == ZEROS
    assert (await _get(stranger, "top", **_today())).json() == []
    assert (await _get(stranger, "dashboard")).json()["debt_total"] == 0
