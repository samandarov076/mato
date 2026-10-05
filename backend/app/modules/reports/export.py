"""Hisobotni Excel fayliga (.xlsx) aylantirish."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Varaq: (nomi, sarlavhalar, qatorlar)
Sheet = tuple[str, list[str], list[list[Any]]]


# Ro'yxat eksportida bitta fayldagi eng ko'p qator (qoldiqda — SKU)
EXPORT_ROWS = 5000
EXPORT_SKUS = 2000


def local_time(moment: datetime | None, timezone: str) -> str:
    """Vaqt — biznes vaqt mintaqasida, o'qishga qulay: "03.10.2026 21:56"."""
    return moment.astimezone(ZoneInfo(timezone)).strftime("%d.%m.%Y %H:%M") if moment else ""


def number(value: Decimal | None) -> float | None:
    """Metr — Excel'da son bo'lib tursin (matn emas): yig'indi va saralash ishlaydi."""
    return None if value is None else float(value)


def workbook(sheets: list[Sheet]) -> bytes:
    """Har varaqda: qalin sarlavha qatori (aylantirganda joyida turadi), ostida ma'lumot."""
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sheets:
        sheet = book.create_sheet(title[:31])  # Excel varaq nomi 31 belgidan oshmaydi
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            longest = max([len(str(header)), *(len(str(row[index - 1])) for row in rows)])
            sheet.column_dimensions[get_column_letter(index)].width = min(longest + 3, 50)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def download(name: str, sheets: list[Sheet]) -> Response:
    """Brauzer faylni yuklab olishi uchun javob. `name` — lotin harflarida, kengaytmasiz."""
    return Response(
        content=workbook(sheets),
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'},
    )
